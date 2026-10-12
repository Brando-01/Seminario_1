"""Create the updated oral/code guide ONLY from a fully audited thesis run."""
import argparse
import json
from pathlib import Path
import numpy as np
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

parser = argparse.ArgumentParser()
parser.add_argument('--run', required=True)
args = parser.parse_args()
root = Path(__file__).resolve().parents[1]
run = root / 'outputs' / args.run
def read(path): return json.loads(path.read_text(encoding='utf-8'))
protocol = read(run / 'protocol.json')
audit = read(run / 'audit.json')
shap_audit = read(run / 'audit_shap.json')
if not all(protocol[key] for key in ('train_complete', 'test_evaluated', 'shap_complete')) or not audit['passed'] or not shap_audit['passed']:
    raise ValueError('The updated guide requires complete training, evaluation, SHAP and audit')
rows = read(run / 'results.json')
summaries = {d: read(run / d / 'prepared' / 'metrics' / 'prepare_summary.json') for d in protocol['datasets']}
shap_summaries = {d: read(run / d / 'explanations' / 'summary.json') for d in protocol['datasets']}
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.left_margin = section.right_margin = Inches(.8)
section.top_margin = section.bottom_margin = Inches(.7)
section.header_distance = section.footer_distance = Inches(.3)
for name in ('Normal','Title','Subtitle','Heading 1','Heading 2'):
    style = doc.styles[name]; style.font.name = 'Calibri'; style.font.color.rgb = RGBColor(0,0,0)
for style in doc.styles:
    for border in list(style.element.iter(qn('w:pBdr'))): border.getparent().remove(border)
doc.styles['Normal'].font.size = Pt(11)
doc.styles['Normal'].paragraph_format.space_after = Pt(7)
doc.styles['Normal'].paragraph_format.line_spacing = 1.08
doc.styles['Title'].font.size = Pt(25)
doc.styles['Heading 1'].font.size = Pt(18)
doc.styles['Heading 2'].font.size = Pt(12)
for name in ('Heading 1','Heading 2'):
    doc.styles[name].paragraph_format.keep_with_next = True
header = section.header.paragraphs[0]
header.text = 'SEMINARIO DE INVESTIGACIÓN 1   |   PROYECTO Y EXPOSICIÓN DEL CÓDIGO'
header.runs[0].font.size = Pt(8)
footer = section.footer.paragraphs[0]
footer.text = 'Beteta Paucar Brandon Jonathan   ·   12 de octubre de 2026   ·   Página '
footer.runs[0].font.size = Pt(8)
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); footer._p.append(field)
doc.core_properties.author = 'Beteta Paucar Brandon Jonathan'
doc.core_properties.title = 'Guía actualizada de exposición del proyecto GNN Autoencoder'

def p(text): return doc.add_paragraph(text)
def h(text): doc.add_heading(text,2)
def page(title): doc.add_page_break(); doc.add_heading(title,1)
def say(text):
    item = p('Qué decir  «' + text + '»'); item.runs[0].italic = True
def code(text):
    item = p(text)
    for r in item.runs: r.font.name = 'Consolas'; r.font.size = Pt(9)
def table(headers, records, widths):
    t = doc.add_table(rows=1,cols=len(headers)); t.autofit=False
    for col,width in zip(t.columns,widths): col.width=Inches(width)
    for cell,text,width in zip(t.rows[0].cells,headers,widths):
        cell.text=text; cell.width=Inches(width)
        shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'EEEEEE'); cell._tc.get_or_add_tcPr().append(shade)
        for r in cell.paragraphs[0].runs:r.bold=True
    repeat=OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for record in records:
        for cell,text,width in zip(t.add_row().cells,record,widths): cell.text=str(text);cell.width=Inches(width)
    for row in t.rows:
        for ci,cell in enumerate(row.cells):
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            margins=OxmlElement('w:tcMar')
            for edge in ('top','bottom','left','right'):
                node=OxmlElement('w:'+edge);node.set(qn('w:w'),'85');node.set(qn('w:type'),'dxa');margins.append(node)
            cell._tc.get_or_add_tcPr().append(margins)
            for para in cell.paragraphs:
                para.paragraph_format.space_after=Pt(3)
                if ci>0:para.alignment=WD_ALIGN_PARAGRAPH.CENTER
                for r in para.runs:r.font.size=Pt(10)
    borders=OxmlElement('w:tblBorders')
    for edge in ('top','bottom','left','right','insideH','insideV'):
        node=OxmlElement('w:'+edge)
        for k,v in [('val','single'),('sz','4'),('color','D9D9D9')]:node.set(qn('w:'+k),v)
        borders.append(node)
    t._tbl.tblPr.append(borders);p('')
def percent(value): return f'{100*value:.2f}'.replace('.',',')
def mean(dataset,model,key,field='metrics'):
    selected=[r for r in rows if r['dataset']==dataset and r['model']==model]
    return float(np.mean([r[field][key] if field else r[key] for r in selected]))

doc.add_paragraph('Guía actualizada para exponer\nel proyecto GNN Autoencoder','Title')
doc.add_paragraph('Código de los dos escenarios evaluación temporal y explicabilidad SHAP','Subtitle')
p('Beteta Paucar Brandon Jonathan\nSeminario de Investigación 1 · 12 de octubre de 2026')
h('Estado del proyecto ejecutable')
p('La ruta actual implementa los seis objetivos específicos: preparación, grafo causal, cuatro modelos, métricas, costo separado y SHAP. Se ejecutaron los cuatro modelos en ambos escenarios con tres semillas, se verificaron sus evidencias y se generaron explicaciones del híbrido. Esto completa el alcance computacional de estas ventanas; no significa que la tesis o una validación de producción estén concluidas.')
p('Esta guía reemplaza la explicación del avance anterior, donde SHAP y la comparación mejorada aún figuraban como pendientes. Conserva los resultados de corridas anteriores solo como antecedentes; las tablas siguientes pertenecen exclusivamente a la corrida actual.')
p('Corrida de evidencia: ' + args.run + '. Auditoría: ' + str(audit['rows']) + ' resultados, mismas particiones por dataset, checkpoints y umbrales congelados, métricas reproducidas.')
h('Apertura de la exposición')
say('Evalué una arquitectura que combina contexto relacional, representación tabular y reconstrucción de un Autoencoder. Utilicé un pipeline común en IEEE-CIS y PaySim, adapté sus relaciones, comparé cuatro modelos y expliqué predicciones del híbrido mediante SHAP. El resultado científico depende del escenario; la arquitectura no garantiza ganar.')
p('Duración sugerida: 30 a 40 minutos. Presenta los objetivos, los datos, el grafo, cada modelo, el entrenamiento y la evaluación; cierra con SHAP y una demostración de resultados guardados.')

page('Objetivos y evidencias de cumplimiento')
p('El objetivo general es evaluar el híbrido frente a modelos individuales y una referencia bajo condiciones comparables. Una evaluación puede cumplir el objetivo aunque el híbrido no sea el mejor modelo.')
table(['Objetivo específico','Componente','Evidencia'],[
('1 Caracterizar y preprocesar ambos datasets','src/data y preparación','Variables retenidas, faltantes, tiempos y conteos'),
('2 Construir relaciones adaptadas','temporal_graph.py','Grafo de transacciones pasado a presente'),
('3 Implementar cuatro modelos comparables','src/models y thesis_experiment.py','Cuatro modelos por dataset y semilla'),
('4 Comparar AP Precision Recall F1','src/evaluation','Scores, umbrales, métricas y auditoría'),
('5 Evaluar costo computacional','Entrenadores y evaluación','Optimización, validación e inferencia separadas'),
('6 Incorporar explicabilidad SHAP','src/explainability','Valores, ranking de muestra y casos locales'),
],[2.65,1.8,2.35])
h('Lo que debes actualizar en la redacción metodológica')
p('Los cortes temporales reales son explícitos y no siempre equivalen a 70/15/15. Indica las ventanas y sus cantidades exactas. Aclara que PaySim principal excluye todos los saldos y que IEEE utiliza un grafo proxy. Describe SHAP como una explicación de atributos de la transacción condicionada al contexto histórico fijo.')
h('Límites que se mantienen')
p('No se ejecuta toda la data original, no se evalúa transferencia IEEE → PaySim y no se establece superioridad estadística con tres semillas del mismo período. Los nuevos bloques son posteriores al piloto registrado, pero no se puede certificar acceso histórico externo. Es necesario reportar estas condiciones.')
say('La contribución es la implementación y evaluación comparable del enfoque y su explicación. Si un modelo simple gana, lo reporto; no modifico el test para que la hipótesis salga favorable.')

page('Organización y ruta del código')
code('configs/              escenarios base\ndata/raw/             CSV originales\nsrc/data/             carga y preparación\nsrc/graph/            aristas y subgrafos causales\nsrc/models/           baseline AE GNN e híbrido\nsrc/training/         pérdidas optimización y lotes\nsrc/explainability/   SHAP del híbrido real\noutputs/thesis_v1_... evidencias de la corrida')
p('Abre primero thesis_experiment.py: build_thesis_config fija las ventanas y variables; create_run guarda el protocolo; train_run aprende los modelos y fija decisiones; evaluate_run produce las métricas; audit_run comprueba integridad; explain_run ejecuta SHAP.')
p('prepare_experiment en run_experiment.py centraliza carga, división temporal, ajuste del preprocesador, matrices y grafos. Este componente se reutiliza para todos los modelos del escenario.')
table(['Orden de exposición','Código a mostrar'],[
('Carga y variables','src/data/loaders.py'),('Particiones y transformación','temporal_split.py y preprocessing.py'),
('Relaciones','src/graph/temporal_graph.py'),('Arquitecturas','src/models/'),
('Aprendizaje','torch_trainer.py graph_batches.py gnn_batches.py'),
('Prueba y evidencia','thesis_experiment.py y metrics.py'),('Explicabilidad','conditional_shap.py y report.py'),
],[2.35,4.45])
say('Tengo dos escenarios dentro de un mismo proyecto. Las configuraciones y claves relacionales son distintas; los modelos y criterios de evaluación son compartidos. No son dos implementaciones desconectadas.')

page('Datos utilizados y particiones cronológicas')
table(['Escenario','Train','Validación','Test'],[
('IEEE-CIS días','[30,50)','[50,60)','[60,75)'),('PaySim pasos','1–16','17–24','25–72'),
],[2.0,1.6,1.6,1.6])
records=[]
for d in protocol['datasets']:
    records.append([d.upper()]+[f'{summaries[d]["splits"][s]["rows"]:,} operaciones\n{summaries[d]["splits"][s]["fraud"]:,} fraudes' for s in ('train','validation','test')])
table(['Dataset','Train','Validación','Test'],records,[1.0,1.9,1.9,2.0])
p('En IEEE se unen train_transaction y train_identity por TransactionID. La etiqueta isFraud está en el archivo etiquetado. El test de investigación es su partición temporal; no son los archivos test oficiales de Kaggle sin etiqueta.')
p('PaySim contiene transacciones sintéticas, monto, tipo, pasos horarios y cuentas de origen/destino. Las cuentas conectan transacciones y permiten historial estrictamente anterior. Las operaciones de la misma hora no se informan entre sí.')
say('Entreno con el pasado, elijo decisiones con validación y evalúo en un período posterior. No uso una partición aleatoria como prueba principal. Tampoco llamo test nuevo al período que ya utilicé para mejorar el modelo.')
p('La configuración se congeló antes de preparar datos. Se guardaron hashes de los CSV, código y parámetros. Es una comprobación local de reproducibilidad, no una certificación de que estos registros nunca se hayan visto fuera del proyecto.')

page('Preprocesamiento sin ajustar con prueba')
p('Abre TabularPreprocessor. fit aprende columnas válidas, medianas de imputación, medias y desviaciones, frecuencias y categorías únicamente con train. transform aplica esos parámetros a validación y test sin reaprenderlos.')
code('preprocessor.fit(train)\nX_train = preprocessor.transform(train)\nX_val = preprocessor.transform(validation)\nX_test = preprocessor.transform(test)')
p('El fragmento expresa la regla del pipeline. Los numéricos se imputan y escalan; el monto usa transformación logarítmica. La versión enriquecida conserva frecuencias categóricas, añade hasta 16 categorías frecuentes por campo e indicadores de faltantes; limita magnitudes estandarizadas a 10.')
table(['Dataset','Variables transformadas comunes'],[(d.upper(),summaries[d]['model_feature_count']) for d in protocol['datasets']],[2.0,4.8])
h('Exclusiones y disponibilidad de información')
p('La etiqueta nunca es una característica. Se excluyen IDs crudos, tiempo absoluto y claves proxy del vector tabular. En PaySim se excluyen isFlaggedFraud, todos los saldos y las inconsistencias de saldo, pero se conservan historia causal y hora cíclica. No describas el piloto anterior con saldos iniciales como si fuera esta corrida estricta.')
say('Todos los modelos reciben las mismas columnas preparadas dentro del escenario. La diferencia es cómo aprovechan esa información, no que el híbrido reciba por accidente una etiqueta o un dato posterior que los demás no tienen.')
p('Los parámetros aprendidos se guardan en preprocessor.pkl y las columnas se enumeran en prepare_summary.json. La carga de un checkpoint propio sirve para reproducir; no cargues pickles de terceros.')

page('Nodos aristas y contexto histórico')
p('Abre build_temporal_entity_graph. En ambos datasets un nodo es una transacción. IEEE relaciona atributos compartidos de tarjeta, correo o dispositivo: proxies, no identidades confirmadas. PaySim relaciona operaciones que comparten una cuenta de origen o destino; las cuentas son claves, no nodos del grafo actual.')
p('La construcción recorre timestamps ordenados, conserva hasta tres vecinos previos por entidad dentro de 72 horas y agrega previous → current. Los registros del mismo instante permanecen pendientes hasta cambiar el tiempo.')
code('for previous, _ in queue:\n    edges.add((previous, current))')
say('La flecha permite que una transacción actual consulte operaciones pasadas relacionadas. Nunca utilizo una transacción futura para explicar una pasada. La etiqueta de un vecino no se convierte en característica.')
records=[]
for d in protocol['datasets']:
    for s in ('train','validation','test'):
        graph=summaries[d]['graphs'][s]
        records.append([d.upper(),s,graph['context_nodes'],graph['target_nodes'],graph['edges']])
table(['Dataset','Partición','Contexto','Objetivos','Aristas'],records,[1.0,1.25,1.5,1.5,1.55])
p('Contexto incluye operaciones anteriores; objetivos son las operaciones para las que se calcula pérdida o métrica. Un grafo de test grande no implica que todos sus nodos se evalúen como casos de test. Los mini-lotes conservan los ancestros exactos necesarios, no una selección aleatoria de vecinos.')

page('Funcionamiento de los cuatro modelos')
h('Baseline')
p('TabularBaseline aplica HistGradientBoostingClassifier con pesos balanceados por clase. Utiliza características tabulares, no aristas. Es el punto de referencia para justificar la complejidad adicional.')
h('Autoencoder')
p('Autoencoder comprime entrada D → 128 → 16 y reconstruye 16 → 128 → D. Aprende con operaciones legítimas de train y conserva el mejor MSE de validación legítima. Su score independiente es el error absoluto medio: anomalía no es necesariamente una probabilidad calibrada de fraude. Las etiquetas seleccionan legítimas, por lo que no se presenta todo el proceso como totalmente no supervisado.')
h('GraphSAGE independiente')
p('GraphSAGELayer promedia vecinos entrantes y concatena ese promedio con la representación propia. Dos capas aprenden contexto relacional. GNNFraudClassifier también incluye una proyección tabular; muestra el código para no describirlo como una rama exclusivamente relacional.')
h('Híbrido')
p('HybridGNNAutoencoder concatena latente AE de 16, embedding GNN de 128, proyección tabular de 128 y error de reconstrucción de 1: vector de 273. LayerNorm y un clasificador 273 → 128 → 1 generan un logit; sigmoide produce el score.')
code('fused = torch.cat([latent, graph_embedding,\n                   tabular_embedding, reconstruction_error], dim=1)\nlogits = self.classifier(self.fusion_normalization(fused))')
say('La fusión no promedia dos predicciones. Aprende a combinar señales. La compuerta de atención repondera atributos dentro del AE; SHAP se calcula después y es un mecanismo diferente.')

page('Entrenamiento semillas y pérdidas')
p('Las redes supervisadas recorren nodos objetivo en lotes de 4 096; la inferencia usa 8 192. IncomingNeighborhoods recupera contexto causal exacto y reutiliza medias del primer nivel. Las pruebas comprueban que la inferencia por lotes coincide con la del grafo completo.')
p('En el híbrido se preentrena el AE hasta 10 épocas y después se ajusta conjuntamente la red hasta 30. Se utiliza Adam con tasa 0,001, regularización 0,0001, dropout 0,15 y recorte de gradiente. La BCE ponderada clasifica y la reconstrucción de operaciones legítimas aporta un término con peso 0,01.')
p('El peso de fraude procede solo de train y se limita a 20. La parada temprana utiliza paciencia 10 y AP de validación. El AE independiente conserva el mejor MSE. No se elige la mejor época por el resultado de test.')
say('Las semillas 42, 52 y 62 cambian inicialización y orden de entrenamiento, no las particiones. Reporto medias y conservo el detalle para ver estabilidad; tres semillas no son tres períodos independientes.')
h('Congelación y reanudación')
p('Antes de test se guardan pesos, scores de validación y dos umbrales. La auditoría comprueba sus hashes. Si se interrumpe el entrenamiento, se reutilizan modelos congelados de esa corrida. Cambiar código, datos o configuración exige iniciar otra corrida; no se edita el manifiesto para eludir las comprobaciones.')
p('La ruta ejecutada utilizó ' + protocol['device'] + '. Los tiempos se miden en este equipo. Una GPU puede acelerar componentes neuronales compatibles, pero no garantiza mejores métricas ni acelera todas las etapas por igual.')

page('Métricas y lectura de los resultados')
table(['Métrica','Explicación oral'],[
('AP','Evalúa el ordenamiento de alertas considerando Precision y Recall a distintos umbrales; es la principal.'),
('Precision','De mis alertas de fraude, qué proporción realmente es fraude.'),
('Recall','De los fraudes reales, qué proporción detecté.'),
('F1','Equilibrio entre Precision y Recall mediante media armónica.'),
('FP por 1 000','Alertas incorrectas por cada 1 000 transacciones legítimas.'),
],[1.25,5.55])
p('La clave auc_pr guarda average_precision_score. En la exposición se llama AP; no es accuracy ni un área PR integrada de forma trapezoidal. AP utiliza scores, mientras Precision, Recall y F1 dependen de un umbral concreto.')
p('El umbral F1 maximiza F1 de validación. El operativo busca Recall con hasta 10 falsas alertas por 1 000 en validación. Son puntos distintos: no mezcles el Recall operativo con Precision y F1 del otro umbral. El presupuesto puede superarse en test por cambios temporales.')
p('Las tablas muestran porcentajes medios de tres semillas. F1 es el promedio de los F1 de cada corrida, no el F1 calculado con Precision y Recall promedios. Si muestras ±, representa dispersión entre semillas, no certeza estadística automática.')
say('No comparo directamente mi F1 con el de un paper sin comprobar sus particiones, variables, desbalance y umbral. Tener otra cifra no demuestra por sí solo que el código esté mal. Ajustar decisiones con test sí invalida esa comparación.')

page('Resultados de la corrida completa')
p('Porcentajes de test, media de semillas 42, 52 y 62. Cada dataset compara cuatro modelos con las mismas variables y particiones. Estas cifras no se mezclan con V2 ni con el piloto híbrido anterior.')
for d in protocol['datasets']:
    h(d.upper())
    record=[]
    for model in ('baseline','autoencoder','gnn','hybrid'):
        record.append([model]+[percent(mean(d,model,key)) for key in ('auc_pr','precision','recall','f1')]
                      +[f'{mean(d,model,"false_positives_per_1000"):.2f}'.replace('.',',')])
    table(['Modelo','AP','Precision','Recall','F1','FP/1 000'],record,[1.55,.95,1.15,.95,.95,1.25])
    best=max(('baseline','autoencoder','gnn','hybrid'),key=lambda m:mean(d,m,'auc_pr'))
    delta=100*(mean(d,'hybrid','auc_pr')-mean(d,'baseline','auc_pr'))
    p(f'Mayor AP media observada: {best}. Diferencia híbrido menos baseline: {delta:+.2f} puntos porcentuales. El ranking pertenece a esta ventana; no acredita superioridad universal.')
p('FP/1 000 corresponde aquí al umbral F1, no al operativo. Revisa semillas y matrices de confusión antes de concluir. La auditoría confirma coherencia de los archivos y métricas, no ausencia de todo sesgo ni significancia estadística.')
if 'paysim' in protocol['datasets']:
    p('En PaySim, el híbrido en el umbral operativo alcanzó Recall ' + percent(mean('paysim','hybrid','recall','operating_metrics')) + ' % y ' + f'{mean("paysim","hybrid","false_positives_per_1000","operating_metrics"):.2f}'.replace('.',',') + ' FP/1 000. El presupuesto de 10 validado no se mantuvo en test: no se presenta como un sistema listo para producción.')

page('Costo computacional de los cuatro modelos')
p('Segundos medios por semilla. Entrenamiento mide optimización e incluye preentrenamiento AE en el híbrido; validación e inferencia de test se registran separadas.')
record=[]
for d in protocol['datasets']:
    for model in ('baseline','autoencoder','gnn','hybrid'):
        record.append([d.upper(),model]+[f'{mean(d,model,key,None):.2f}'.replace('.',',') for key in ('seconds','validation_seconds','inference_seconds')])
table(['Dataset','Modelo','Train s','Validación s','Inferencia s'],record,[1.0,1.7,1.3,1.4,1.4])
p('Se excluyen lectura de CSV, ajuste general del preprocesador, construcción inicial del grafo, escritura de checkpoints y cálculo SHAP. La inferencia de grafos sí incluye la preparación de medias y subgrafos utilizada para producir sus scores.')
say('La arquitectura debe justificarse por lo que aporta y por lo que cuesta. Un modelo más lento no es necesariamente mejor. Esta tabla describe mi equipo y ventanas; no es una medición bancaria en tiempo real.')
h('Comprobación de reproducibilidad')
p('protocol.json conserva configuración, semillas, versiones y hashes. training_results.json conserva selección de validación. predictions.npz conserva scores, etiquetas y timestamps. audit.json comprueba checkpoints, mismas particiones, umbrales congelados, métricas y matrices de confusión.')

page('Explicabilidad SHAP del híbrido')
p('Abre src/explainability/conditional_shap.py. FixedContextPredictor reconstruye la computación del híbrido real en subgrafos disjuntos y cambia únicamente los atributos del nodo objetivo. GroupCoalitions reúne codificaciones e indicadores del mismo nombre de columna y los reemplaza usando referencias de train.')
p('PermutationExplainer estima contribuciones mediante permutaciones. No utiliza un clasificador sustituto ni interpreta pesos de atención como SHAP. Se verifica tanto la coincidencia con el score guardado como la identidad base más suma de contribuciones igual al score, dentro de tolerancia numérica.')
say('SHAP me permite ver qué grupos de atributos aumentan o reducen la puntuación del híbrido para una transacción. Mantengo fijas las relaciones y la historia. Por eso no digo que expliqué causalmente el fraude ni que medí la importancia de una arista.')
p('La explicación usa ' + str(protocol['explanations']['background_train_rows']) + ' filas de train como fondo, ' + str(protocol['explanations']['permutations']) + ' permutaciones y ' + str(protocol['explanations']['random_sample']) + ' transacciones aleatorias de test. La importancia global es de esa muestra, no de toda la población. Se añaden TP, FP, FN y TN ilustrativos cuando existen; esos extras no forman parte de la media global.')
table(['Dataset','Grupos','Error máximo de aditividad'],[(d.upper(),shap_summaries[d]['group_count'],f'{shap_summaries[d]["max_additivity_error"]:.3g}') for d in protocol['datasets']],[2.0,1.3,3.5])
p('Las unidades son probabilidad. Contribuciones positivas aumentan el score; negativas lo reducen. Se explican variables transformadas, no unidades monetarias originales. El fondo y las permutaciones influyen. Las variables ingenierizadas, como monto y log-monto, pueden ser grupos separados: las perturbaciones no garantizan transacciones físicamente plausibles.')
p('Referencia del método: shap.readthedocs.io/en/stable/generated/shap.PermutationExplainer.html. En la carpeta explanations se guardan valores, resumen, ranking y gráficos locales.')

for d in protocol['datasets']:
    page('Interpretación de SHAP en ' + d.upper())
    p('Media de las contribuciones SHAP absolutas de 16 transacciones de test elegidas uniformemente. Los casos ilustrativos adicionales no se incluyen en este ranking; no representa toda la población.')
    doc.add_picture(str(run / d / 'explanations' / 'global_importance.png'),width=Inches(6.45))
    top=shap_summaries[d]['top_features'][:3]
    p('Mayor contribución absoluta media en esta muestra: ' + ', '.join(item['feature'] for item in top) + '. La magnitud no muestra por sí sola la dirección del efecto; abre un caso local para ver qué sube o baja su score.')
    p('En el gráfico local, la base es el score medio al reemplazar los atributos con el fondo de train y mantener fijo el contexto. Rojo aumenta el score, azul lo reduce; base más contribuciones llega al score explicado. La base puede cambiar de una transacción a otra.')
    say('Este ranking pertenece a una muestra pequeña y al fondo elegido. No afirmo que estos campos sean causas reales de fraude. En los ejemplos locales comparo una alerta correcta, una falsa alerta y un fraude no detectado para entender el comportamiento del modelo.')

page('Demostración y cierre de la exposición')
p('Dentro del proyecto, con (.venv) activo, muestra la corrida guardada sin volver a entrenar:')
code('python show_results.py --thesis-run ' + args.run + ' --details\npython thesis_experiment.py --stage audit --run ' + args.run)
p('La primera orden muestra métricas y costos; la segunda comprueba hashes y reproduce las métricas guardadas. No cambian los umbrales ni reinician el aprendizaje.')
h('Panel y explicaciones')
code('python dashboard_app.py --no-prepare')
p('Abre http://127.0.0.1:8765 y selecciona la corrida thesis_v1. El selector actualiza tablas y curvas; el detalle SHAP muestra ranking y casos. La animación conserva una caché independiente con su fuente identificada. --rebuild-cache genera ejemplos de la corrida thesis_v1 completa más reciente; cambiar el selector a otra corrida no reconstruye esas animaciones. Si el servidor anterior ya está abierto, usa otra terminal con --port 8766 para cargar el código actualizado.')
h('Cómo reproducir todo fuera de la exposición')
code('python -m pip install -r requirements.txt\npython -m pytest tests -q\npython thesis_experiment.py --stage all')
p('La última orden crea otra corrida y puede tardar. No es necesaria para presentar resultados guardados. Copia los CSV por separado al cambiar de computadora; .venv y outputs no se incluyen automáticamente en Git.')
h('Cierre de treinta segundos')
say('Implementé un pipeline temporal común, relaciones adaptadas a cada dataset, cuatro modelos comparables, métricas de desbalance y medición de costo. Añadí SHAP del híbrido con alcance explícito y guardé evidencias auditables. Mis conclusiones se basan en la corrida y sus límites; una arquitectura más compleja debe demostrar su utilidad, no asumirla.')
p('El código ejecutable cumple el alcance de los seis objetivos para estas ventanas. La redacción de la ficha y el artículo debe alinearse con los cortes reales, el escenario estricto, el alcance SHAP y las conclusiones observadas. La publicación en GitHub requiere una decisión posterior; no se hizo push automático.')

for paragraph in doc.paragraphs:
    for border in list(paragraph._p.iter(qn('w:pBdr'))):border.getparent().remove(border)
output=root / 'docs' / 'Guia_exposicion_proyecto_completo.docx'
doc.save(output)
print(output)
